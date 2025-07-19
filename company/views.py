from django.shortcuts import render
from django.http import HttpResponse
from .models import *
from django.utils import timezone
from .serializers import *
from rest_framework import viewsets
from rest_framework.decorators import action
from django.shortcuts import get_object_or_404
from rest_framework import status
from rest_framework.authentication import TokenAuthentication
from rest_framework import generics
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from django.contrib.auth.models import User
from rest_framework.permissions import AllowAny
from rest_framework_simplejwt.tokens import RefreshToken
from django.contrib.auth import authenticate
from rest_framework import permissions
from django.db import transaction
import pandas as pd
from rest_framework.exceptions import ValidationError
import os
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import JsonResponse
from rest_framework.views import APIView
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from .permissions import *
from django.contrib.auth.password_validation import validate_password


#+=========================================================================================================================
#============================                   LOGIN                              ======================================================================
#======================================================================================================== 



class LoginView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        username = request.data.get("username")
        password = request.data.get("password")

        if not username or not password:
            return Response({"error": "Username and password are required", "status": 400})

        user = authenticate(username=username, password=password)

        if user is not None:
            user_role = "owner" if hasattr(user, 'owner') else "user"

            # ✅ Enforce check: If Owner exists, ensure they have at least one company
            if user_role == "owner":
                owner = user.owner
                if not owner.company_set.exists():  # or replace with your relation field
                    return Response({"error": "Owner has no associated company", "status": 403})

            refresh = RefreshToken.for_user(user)
            return Response({
                "refresh": str(refresh),
                "access": str(refresh.access_token),
                "username": user.username,
                "user_role": user_role,
                "status": 200
            })

        return Response({"error": "Invalid credentials", "status": 400})



#+=========================================================================================================================
#============================                   LOGINOUT                             ======================================================================
#======================================================================================================== 



class LogoutView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        
        refresh_token = request.data["refresh"]
        if not refresh_token:
             return Response({"error": "Refresh token is required", "status": 400})
        try:
            token = RefreshToken(refresh_token)
            token.blacklist()
            return Response({"msg": "Logout successful", "status": 200})
        
        except Exception as e:
            return Response({"error": str(e), "status": 400})


#+=========================================================================================================================
#============================               OWNER                              ======================================================================
#======================================================================================================== 


class CreateOwnerView(APIView):
    # You can make it public or protected — up to you
    #permission_classes = [AllowAny]
    # @swagger_auto_schema(request_body=ClientSerializer)
    permission_classes = [AllowAny]
    def post(self, request):
        data = request.data
        username = data.get('username')
        email = data.get('email')
        password = data.get('password')
        phone_no = data.get('phone_no')
        type_of_company = data.get('type_of_company')  # 'basic' or 'premium'

        if not all([username, email, password, phone_no, type_of_company]):
            return Response({"error": "All fields are required."}, status=400)

        if User.objects.filter(username=username).exists():
            return Response({"error": "Username already exists"}, status=400)

        if User.objects.filter(email=email).exists():
            return Response({"error": "Email already exists"}, status=400)

        if type_of_company not in ['basic', 'premium']:
            return Response({"error": "type_of_company must be 'basic' or 'premium'"}, status=400)

        try:
            with transaction.atomic(): # will roleback if any error occurs
                # Create the user and owner in a single transaction
                user = User.objects.create_user(
                    username=username,
                    email=email,
                    password=password
                )

                owner = Owner.objects.create(
                    user=user,
                    phone_no=phone_no,
                    type_of_company=type_of_company
                )


            return Response({
                "message": "Owner created successfully",
                "user_id": user.id,
                "owner_id": owner.id
                , "status": 201
            })

        except Exception as e:
            return Response({"error": str(e), "status": 500})

class OwnerDetailView(APIView):
    permission_classes = [IsAuthenticated,IsOwner]

    def get(self, request, pk):
        try:
            owner = Owner.objects.get(pk=pk)
        except Owner.DoesNotExist:
            return Response({"error": "Owner not found" , "status": 404})

        serializer = OwnerSerializer(owner)
        return Response({"info":serializer.data, "status":200})
    
class ListOwnersView(APIView):
   # permission_classes = [AllowAny]
    permission_classes = [IsOwner,IsAuthenticated]
    def get(self, request):
        owners = Owner.objects.filter(deleted=False)
        serializer = OwnerSerializer(owners, many=True)
        return Response(serializer.data, status=200)

class RetrieveOwnerView(APIView):
   # permission_classes = [AllowAny]
    permission_classes = [IsAuthenticated,IsOwner]
    def get(self, request, pk):
        owner = get_object_or_404(Owner, pk=pk,deleted=False)
        serializer = OwnerSerializer(owner)
        return Response(serializer.data, status=200)

class UpdateOwnerView(APIView):
    permission_classes=[IsAuthenticated,IsOwner]
    def put(self, request, pk):
        try:
            owner = Owner.objects.get(pk=pk, deleted=False)
        except Owner.DoesNotExist:
            return Response({"status": 404, "message": "Owner not found"})

        serializer = OwnerSerializer(owner, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response({"status": 200, "message": "Owner updated successfully", "data": serializer.data})
        return Response({"status": 400, "errors": serializer.errors})

class DeleteOwnerView(APIView):
    permission_classes=[IsAuthenticated,IsOwner]
    def delete(self, request, pk):
        try:
            owner = Owner.objects.get(pk=pk, deleted=False)
        except Owner.DoesNotExist:
            return Response({"status": 404, "message": "Owner not found"})

        owner.deleted = True
        owner.save()
        return Response({"status": 200, "message": "Owner soft deleted successfully"})






#+=========================================================================================================================
#============================               COMPANY                              ======================================================================
#======================================================================================================== 




class CreateCompanyView(APIView):
    permission_classes = [IsAuthenticated,IsOwner]

    def post(self, request):
        data = request.data
        owner_id = data.get('owner_id')
        company_name = data.get('company_name')
        phone_no = data.get('phone_no')
        gst = data.get('gst')
        address = data.get('address')
        type_of_company = data.get('type_of_company')

        # Bank details from request
        bank_name = data.get('bank_name')
        account_holder_name = data.get('account_holder_name')
        account_no = data.get('account_no')
        ifsc_code = data.get('ifsc_code')
        bank_address = data.get('bank_address')
        branch = data.get('branch')
        ad_code = data.get('ad_code')
        swift_code = data.get('swift_code')
        opening_balance = data.get('opening_balance', 0.0)
        as_on = data.get('as_on')

        # Validate company fields
        if not all([owner_id, company_name, phone_no, gst, address]):
            return Response({"error": "Missing required company fields", "status": 400}, status=400)
        # Validate bank fields
        if not all([bank_name, account_holder_name, account_no, ifsc_code, bank_address, branch, opening_balance, as_on]):
            return Response({"error": "Missing required bank fields", "status": 400}, status=400)

        try:
            owner = Owner.objects.get(id=owner_id)
        except Owner.DoesNotExist:
            return Response({"error": "Owner not found", "status": 404}, status=404)

        limit = 3 if owner.type_of_company == 'basic' else 5
        current = Company.objects.filter(owner=owner).count()

        if current >= limit:
            return Response({
                "error": f"{owner.type_of_company} plan allows only {limit} companies",
                "status": 400
            }, status=400)

        try:
            with transaction.atomic():
                company = Company.objects.create(
                    owner=owner,
                    company_name=company_name,
                    phone_no=phone_no,
                    gst=gst,
                    address=address,
                    type_of_company=type_of_company
                )
                initial_cash = float(data.get("initial_cash", 0.0))  # default to 0.0

                # After creating the company
                if initial_cash > 0:
                    CashLedger.objects.create(
                        company=company,
                        amount=initial_cash,
                        transaction_type='inflow',
                        description="Initial cash in hand"
                    )

                Bank.objects.create(
                    company=company,
                    bank_name=bank_name,
                    account_holder_name=account_holder_name,
                    account_no=account_no,
                    ifsc_code=ifsc_code,
                    address=bank_address,
                    branch=branch,
                    ad_code=ad_code,
                    swift_code=swift_code,
                    opening_balance=opening_balance,
                    as_on=as_on
                )

            return Response({
                "msg": "Company and Bank created successfully",
                "company_id": company.id
            }, status=201)

        except Exception as e:
            return Response({"error": str(e), "status": 500})

class ListCompaniesView(APIView):
   # permission_classes = [AllowAny]
    permission_classes = [IsAuthenticated,IsOwner]
    def get(self, request):
        companies = Company.objects.filter(deleted=False)
        serializer = CompanySerializer(companies, many=True)
        return Response(serializer.data, status=200)

class RetrieveCompanyView(APIView):
   # permission_classes = [AllowAny]
    permission_classes = [IsAuthenticated,IsOwner] 
    def get(self, request, pk):    
        company = get_object_or_404(Company, pk=pk,deleted=False)
        serializer = CompanySerializer(company)
        return Response(serializer.data, status=200)

class UpdateCompanyView(APIView):
    permission_classes=[IsAuthenticated,IsOwner]
    def put(self, request, pk):
        try:
            company = Company.objects.get(pk=pk, deleted=False)
        except Company.DoesNotExist:
            return Response({"status": 404, "message": "Company not found"})

        serializer = CompanySerializer(company, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response({"status": 200, "message": "Company updated successfully", "data": serializer.data})
        return Response({"status": 400, "errors": serializer.errors})

class DeleteCompanyView(APIView):
    permission_classes=[IsAuthenticated,IsOwner]
    def delete(self, request, pk):
        try:
            company = Company.objects.get(pk=pk, deleted=False)
        except Company.DoesNotExist:
            return Response({"status": 404, "message": "Company not found"})

        company.deleted = True
        company.save()
        return Response({"status": 200, "message": "Company soft deleted successfully"})









#+=========================================================================================================================
#============================               CLIENT                              ======================================================================
#======================================================================================================== 






class CreateClientView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name = "client"
    def post(self, request):
        data = request.data

        required_fields = ['client_type', 'client_name', 'mobile_number', 'email', 'gstin', 'state', 'billing_address']
        if not all(data.get(field) for field in required_fields):
            return Response({"error": "Missing required fields", "status": 400}, status=400)

        try:
            client = Client.objects.create(
                client_type=data.get('client_type'),
                client_name=data.get('client_name'),
                mobile_number=data.get('mobile_number'),
                email=data.get('email'),
                gstin=data.get('gstin'),
                pan=data.get('pan'),
                state=data.get('state'),
                billing_address=data.get('billing_address'),
                billing_address_line2=data.get('billing_address_line2'),
                shipping_address=data.get('shipping_address'),
                pincode_special_economic_zone=data.get('pincode_special_economic_zone', False),
                city=data.get('city'),
                credit_period=data.get('credit_period', 0),
                credit_limit=data.get('credit_limit', 0.0),
                opening_balance=data.get('opening_balance', 0.0),
                other_currency=data.get('other_currency', False),
                check_discount=data.get('check_discount', False),
                enable_multiple_address=data.get('enable_multiple_address', False)
            )
            return Response({
                "msg": "Client created successfully",
                "client_id": client.id
            }, status=201)

        except Exception as e:
            return Response({"error": str(e), "status": 500}, status=500)

class RetrieveClientView(APIView):
   # permission_classes = [AllowAny]
    module_name = "client"
    is_view_specific = True 
   
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    def get(self, request, pk):
        client = get_object_or_404(Client, pk=pk,deleted=False)
        serializer = ClientSerializer(client)
        return Response(serializer.data, status=200)        

class ListClientsView(APIView):
  #  permission_classes = [AllowAny]
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name = "client"
    def get(self, request):
        clients = Client.objects.filter(deleted=False)
        serializer = ClientSerializer(clients, many=True)
        return Response(serializer.data, status=200)

class UpdateClientView(APIView):
    permission_classes=[IsAuthenticated,OwnerOrEmployee]
    module_name = "client"
    def put(self, request, pk):
        try:
            client = Client.objects.get(pk=pk, deleted=False)
        except Client.DoesNotExist:
            return Response({"status": 404, "message": "Client not found"})

        serializer = ClientSerializer(client, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response({"status": 200, "message": "Client updated successfully", "data": serializer.data})
        return Response({"status": 400, "errors": serializer.errors})

class DeleteClientView(APIView):
    permission_classes=[IsAuthenticated,OwnerOrEmployee]
    module_name = "client"
    def delete(self, request, pk):
        try:
            client = Client.objects.get(pk=pk, deleted=False)
        except Client.DoesNotExist:
            return Response({"status": 404, "message": "Client not found"})

        client.deleted = True
        client.save()
        return Response({"status": 200, "message": "Client soft deleted successfully"})











#+=========================================================================================================================
#============================               ITEMS                              ======================================================================
#======================================================================================================== 





class CreateItemView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "item"
    
    def post(self, request,company_id):

        data = request.data

        company_id = data.get('company_id')
        item_name = data.get('item_name')
        item_code = data.get('item_code')
        unit_id = data.get('unit')
        quantity_raw = data.get('quantity')
        description = data.get('description')
        tax_type_id = data.get('tax_type')
        tax = data.get('tax')
        price = data.get('price')
        selling_price = data.get('selling_price', price)

        # --- Validate required fields ---
        if not all([company_id, item_name, item_code, unit_id, description, tax_type_id, price]):
            return Response({"error": "Missing required fields", "status": 400}, status=400)

        # --- Validate numeric fields ---
        try:
            quantity = float(quantity_raw)
            if quantity <= 0:
                return Response({"error": f"Quantity must be > 0. Got {quantity}"}, status=400)
        except (ValueError, TypeError):
            return Response({"error": "Quantity must be a valid number"}, status=400)

        try:
            price = float(price)
        except (ValueError, TypeError):
            return Response({"error": "Price must be a valid number"}, status=400)

        try:
            selling_price = float(selling_price or price)
        except (ValueError, TypeError):
            return Response({"error": "Selling price must be a valid number"}, status=400)

        # --- Get related objects ---
        try:
            company = Company.objects.get(id=company_id)
        except Company.DoesNotExist:
            return Response({"error": "Company not found", "status": 404}, status=404)

        if Item.objects.filter(item_code=item_code).exists():
            return Response({"error": "Item code must be unique", "status": 400}, status=400)

        try:
            unit = Unit.objects.get(id=unit_id)
        except Unit.DoesNotExist:
            return Response({"error": "Unit not found", "status": 404}, status=404)

        try:
            tax_type = TaxType.objects.get(id=tax_type_id)
        except TaxType.DoesNotExist:
            return Response({"error": "TaxType not found", "status": 404}, status=404)

        # --- Tax and price calculation ---
        if tax_type.code == '1':  # assuming '1' is the code for "with tax"
            if tax is None:
                return Response({"error": "Tax is required for 'withtax' items", "status": 400}, status=400)
            try:
                tax = float(tax)
            except ValueError:
                return Response({"error": "Invalid tax value", "status": 400}, status=400)

            price = round(price * (1 + tax / 100), 2)
            selling_price = round(selling_price * (1 + tax / 100), 2)
        else:
            tax = 0
            price = round(price, 2)
            selling_price = round(selling_price, 2)

        # --- Create Item ---
        try:
            item = Item.objects.create(
                company=company,
                item_name=item_name,
                item_code=item_code,
                unit=unit,
                quantity=quantity,
                description=description,
                tax_type=tax_type,
                tax=tax,
                price=price,
                selling_price=selling_price
            )

            return Response({
                "msg": "Item created successfully",
                "item_id": item.id
            }, status=201)

        except Exception as e:
            return Response({"error": str(e), "status": 500}, status=500)

class ListItemsView(APIView):
   # permission_classes = [AllowAny]
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "item"
    def get(self, request,company_id):
        items = Item.objects.filter(deleted=False)
        serializer = ItemSerializer(items, many=True)
        return Response(serializer.data, status=200)

class RetrieveItemView(APIView):
   # permission_classes = [AllowAny]
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "item"
    is_view_specific = True 
    def get(self, request,company_id, pk):
        item = get_object_or_404(Item,pk=pk,deleted=False)
        serializer = ItemSerializer(item)
        return Response(serializer.data, status=200)

class UpdateItemView(APIView):
    permission_classes=[IsAuthenticated, OwnerOrEmployee]
    module_name = "item"
    def put(self, request, company_id,pk):
        try:
            item = Item.objects.get(pk=pk,deleted=False)
        except Item.DoesNotExist:
            return Response({"status": 404, "message": "Item not found"})

        serializer = ItemSerializer(item, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response({"status": 200, "message": "Item updated successfully", "data": serializer.data})
        return Response({"status": 400, "errors": serializer.errors})

class DeleteItemView(APIView):
   # permission_classes = [AllowAny]
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "item"
    def delete(self, request,company_id,pk):
        item = get_object_or_404(Item,pk=pk)
        item.delete()
        return Response({"msg": "Item deleted", "status": 200})









#+=========================================================================================================================
#============================               SALES_INVOICE                              ======================================================================
#======================================================================================================== 


class CreateSalesInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="sales-invoice"
    def post(self, request,company_id):
        serializer = SalesInvoiceSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        data = serializer.validated_data
        items_data = request.data.get("items", [])

        # Additional required fields
        payment_mode_id = request.data.get("payment_mode")
        payment_status_id = request.data.get("payment_status")
        payment_type_id = request.data.get("payment_type")
        received_amt = request.data.get("received_amt", 0.0)

        if not all([payment_mode_id, payment_status_id, payment_type_id]):
            return Response({"error": "Missing payment_mode, payment_status, or payment_type"}, status=400)

        try:
            with transaction.atomic():
                payment_mode = PaymentMode.objects.get(id=payment_mode_id)
                payment_status = PaymentStatus.objects.get(id=payment_status_id)
                payment_type = PaymentType.objects.get(id=payment_type_id)

                invoice = SalesInvoice.objects.create(
                    company=data["company"],
                    client=data["client"],
                    invoice_number=data["invoice_number"],
                    final_discount_applicable=data.get("final_discount_applicable", False),
                    final_discount=data.get("final_discount", 0),
                    payment_mode=payment_mode,
                    payment_status=payment_status,
                    payment_type=payment_type,
                    received_amt=received_amt,
                )

                subtotal = 0
                for item_data in items_data:
                    item_id = item_data.get("item")
                    quantity = item_data.get("quantity")
                    discount_applicable = item_data.get("discount_applicable", False)
                    discount = item_data.get("discount", 0)

                    if not all([item_id, quantity]):
                        raise ValueError("Item ID and quantity are required")

                    item = Item.objects.get(id=item_id)

                    if item.quantity < quantity:
                        raise ValueError(
                            f"Not enough stock for item: {item.item_name}. "
                            f"Available: {item.quantity}, Requested: {quantity}"
                        )

                    item.quantity -= quantity
                    item.save()

                    selling_price = item.selling_price or 0
                    line_total = quantity * selling_price
                    if discount_applicable and discount > 0:
                        line_total -= line_total * (discount / 100)

                    line_total = round(line_total, 2)
                    subtotal += line_total

                    SalesInvoiceItem.objects.create(
                        invoice=invoice,
                        item=item,
                        quantity=quantity,
                        discount_applicable=discount_applicable,
                        discount=discount,
                        line_total=line_total
                    )

                invoice.total_price = invoice.apply_final_discount(subtotal)
                invoice.save(update_fields=["total_price"])

                return Response(SalesInvoiceSerializer(invoice).data, status=201)

        except Exception as e:
            return Response({"error": str(e)}, status=500)

class ListSalesInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="sales-invoice"

    def get(self, request,company_id):
        invoices = SalesInvoice.objects.filter(company_id=company_id,deleted=False)
        serializer = SalesInvoiceSerializer(invoices, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

class SoftDeleteSalesInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="sales-invoice"

    def delete(self, request,company_id, pk):
        try:
            invoice = SalesInvoice.objects.get(pk=pk, company_id=company_id,deleted=False)
            items = invoice.items.all()

            for item_entry in items:
                item = item_entry.item
                item.quantity += item_entry.quantity
                item.save()

            invoice.deleted = True  # ✅ fix field name
            invoice.save(update_fields=['deleted'])

            return Response({"detail": "Invoice soft-deleted"}, status=status.HTTP_200_OK)

        except SalesInvoice.DoesNotExist:
            return Response({"error": "Invoice not found or already deleted"}, status=status.HTTP_404_NOT_FOUND)
    
class RetrieveSalesInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="sales-invoice"
    module_name = "item"

    def get(self, request, company_id,pk):
        try:
            invoice = SalesInvoice.objects.get(pk=pk, company_id=company_id,deleted=False)
            serializer = SalesInvoiceSerializer(invoice)
            return Response(serializer.data, status=200)
        except SalesInvoice.DoesNotExist:
            return Response({"error": "Sales invoice not found"}, status=404)

class UpdateSalesInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="sales-invoice"

    def put(self, request, company_id,pk):
        try:
            invoice = SalesInvoice.objects.get(pk=pk, company_id=company_id,deleted=False)
        except SalesInvoice.DoesNotExist:
            return Response({"error": "Sales invoice not found"}, status=404)

        # Exclude current invoice from unique check
        invoice_number = request.data.get("invoice_number")
        if invoice_number and SalesInvoice.objects.exclude(pk=pk).filter(invoice_number=invoice_number).exists():
            return Response({"invoice_number": ["sales invoice with this invoice number already exists."]}, status=400)

        serializer = SalesInvoiceSerializer(invoice, data=request.data, partial=True)
        if not serializer.is_valid():
            return Response(serializer.errors, status=400)

        items_data = request.data.get("items", [])
        try:
            with transaction.atomic():
                # Restore stock for old items
                for item_entry in invoice.items.all():
                    item = item_entry.item
                    item.quantity += item_entry.quantity
                    item.save()
                invoice.items.all().delete()

                subtotal = 0   
                for item_data in items_data:
                    item_id = item_data.get("item")
                    quantity = item_data.get("quantity")
                    discount_applicable = item_data.get("discount_applicable", False)
                    discount = item_data.get("discount", 0)

                    if not all([item_id, quantity]):
                        raise ValidationError("Item ID and quantity are required")

                    item = Item.objects.get(id=item_id)
                    if item.quantity < quantity:
                        raise ValidationError(
                            f"Not enough stock for item: {item.item_name}. "
                            f"Available: {item.quantity}, Requested: {quantity}"
                        )
                    item.quantity -= quantity
                    item.save()

                    selling_price = item.selling_price or 0
                    line_total = quantity * selling_price
                    if discount_applicable and discount > 0:
                        line_total -= line_total * (discount / 100)
                    line_total = round(line_total, 2)
                    subtotal += line_total

                    SalesInvoiceItem.objects.create(
                        invoice=invoice,
                        item=item,
                        quantity=quantity,
                        discount_applicable=discount_applicable,
                        discount=discount,
                        line_total=line_total
                    )

                serializer.save(total_price=invoice.apply_final_discount(subtotal))
                return Response(SalesInvoiceSerializer(invoice).data, status=200)
        except Exception as e:
            return Response({"error": str(e)}, status=400)







#+=========================================================================================================================
#============================             PURCHASE_INVOICE                              ======================================================================
#======================================================================================================== 




class CreatePurchaseInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="purchase-invoice"

    def post(self, request,company_id):
        data = request.data
        company_id = data.get('company')
        supplier_name = data.get('supplier_name')
        invoice_number = data.get('invoice_number')
        items = data.get('items', [])

        payment_mode_id = data.get('payment_mode')
        payment_status_id = data.get('payment_status')
        payment_type_id = data.get('payment_type')
        print("Request Data:", data)
        print("company_id:", company_id)
        print("items:", items)
        print("payment_mode_id:", payment_mode_id)
        print("payment_status_id:", payment_status_id)
        print("payment_type_id:", payment_type_id)

        # ✅ Validate required fields
        required_fields = [company_id, supplier_name, invoice_number, items]
        if any(field is None for field in required_fields):
            return Response({"error": "Missing required invoice fields"}, status=400)

        if payment_mode_id is None or payment_status_id is None or payment_type_id is None:
            return Response({"error": "Missing required payment fields"}, status=400)

        try:
            company = Company.objects.get(id=company_id)
        except Company.DoesNotExist:
            return Response({"error": "Invalid company ID"}, status=404)

        if PurchaseInvoice.objects.filter(invoice_number=invoice_number, company=company).exists():
            return Response({"error": "Invoice number must be unique within the company"}, status=400)

        try:
            payment_mode = PaymentMode.objects.get(id=payment_mode_id)
            payment_status = PaymentStatus.objects.get(id=payment_status_id)
            payment_type = PaymentType.objects.get(id=payment_type_id)
        except Exception as e:
            return Response({"error": f"Invalid payment field: {str(e)}"}, status=404)

        # ✅ Begin invoice creation
        try:
            with transaction.atomic():
                invoice = PurchaseInvoice.objects.create(
                    company=company,
                    supplier_name=supplier_name,
                    invoice_number=invoice_number,
                    payment_mode=payment_mode,
                    payment_status=payment_status,
                    payment_type=payment_type,
                    paid_amt=0  # always initialize as 0
                )

                total_price = 0

                for item_data in items:
                    item_name = item_data.get('item_name')
                    unit_id = item_data.get('unit')
                    quantity = item_data.get('quantity')
                    cost_price = item_data.get('cost_price')

                    if not all([item_name, unit_id, quantity, cost_price]):
                        raise ValueError(f"Missing required item fields for item '{item_name or 'Unnamed'}'")

                    try:
                        unit = Unit.objects.get(id=unit_id)
                    except Unit.DoesNotExist:
                        raise ValueError(f"Invalid unit ID '{unit_id}' for item '{item_name}'")

                    line_total = round(quantity * cost_price, 2)
                    total_price += line_total

                    # Check if item with same price exists
                    existing_items = Item.objects.filter(item_name=item_name, company=company)
                    matched_item = None
                    for item in existing_items:
                        if round(item.price, 2) == round(cost_price, 2):
                            matched_item = item
                            break

                    if matched_item:
                        matched_item.quantity += quantity
                        matched_item.save()
                    else:
                        matched_item = Item.objects.create(
                            company=company,
                            item_name=item_name,
                            item_code=f"{item_name[:3].upper()}_{Item.objects.count() + 1}",
                            quantity=quantity,
                            unit=unit,
                            description="Auto-created from Purchase Invoice",
                            tax_type=None,
                            tax=None,
                            price=cost_price,
                            selling_price=round(cost_price * 1.1, 2)
                        )

                    PurchaseInvoiceItem.objects.create(
                        invoice=invoice,
                        item_name=item_name,
                        unit=unit,
                        quantity=quantity,
                        cost_price=cost_price,
                        line_total=line_total
                    )

                invoice.total_price = round(total_price, 2)
                invoice.save(update_fields=["total_price"])

                return Response({
                    "msg": "Purchase invoice created successfully",
                    "invoice_id": invoice.id,
                    "total_price": invoice.total_price
                }, status=201)

        except ValueError as ve:
            invoice.delete()
            return Response({"error": str(ve)}, status=400)
        except Exception as e:
            return Response({"error": f"Unexpected error: {str(e)}"}, status=500)

class UpdatePurchaseInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="purchase-invoice"

    def put(self, request, company_id,pk):
        try:
            invoice = PurchaseInvoice.objects.get(pk=pk, company_id=company_id,deleted=False)
        except PurchaseInvoice.DoesNotExist:
            return Response({"error": "Purchase invoice not found"}, status=404)

        data = request.data
        items = data.get('items', [])

        if not items:
            return Response({"error": "At least one item is required"}, status=400)

        try:
            with transaction.atomic():
                # STEP 1: Rollback old stock using exact item match (item_name + unit + cost_price)
                for pi_item in invoice.items.all():
                    try:
                        matched_item = Item.objects.get(
                            item_name=pi_item.item_name,
                            unit=pi_item.unit,
                            company=invoice.company,
                            price=pi_item.cost_price
                        )
                    except Item.DoesNotExist:
                        raise ValidationError(
                            f"No matching item found for rollback: {pi_item.item_name} at price {pi_item.cost_price}"
                        )

                    if matched_item.quantity < pi_item.quantity:
                        raise ValidationError(
                            f"Cannot update invoice: Not enough stock to remove for item '{matched_item.item_name}'. "
                            f"Current stock: {matched_item.quantity}, to remove: {pi_item.quantity}"
                        )

                    matched_item.quantity -= pi_item.quantity
                    matched_item.save()

                invoice.items.all().delete()

                # STEP 2: Add new items
                total_price = 0
                for item_data in items:
                    item_name = item_data.get('item_name')
                    unit_id = item_data.get('unit')
                    quantity = item_data.get('quantity')
                    cost_price = item_data.get('cost_price')

                    if not item_name or not unit_id or quantity in [None, "", 0] or cost_price in [None, "", 0]:
                        raise ValidationError("All item fields are required and must be greater than zero.")

                    try:
                        quantity = float(quantity)
                        cost_price = float(cost_price)
                    except Exception:
                        raise ValidationError("Quantity and cost price must be valid numbers.")

                    if quantity <= 0:
                        raise ValidationError(f"Quantity must be greater than zero for item: {item_name}")
                    if cost_price <= 0:
                        raise ValidationError(f"Cost price must be greater than zero for item: {item_name}")

                    try:
                        unit = Unit.objects.get(id=unit_id)
                    except Unit.DoesNotExist:
                        raise ValidationError(f"Invalid unit ID '{unit_id}' for item '{item_name}'.")

                    line_total = round(quantity * cost_price, 2)
                    total_price += line_total

                    # STEP 2.1: Reuse existing item if same name, company, unit and price match
                    existing_items = Item.objects.filter(
                        item_name=item_name,
                        company=invoice.company,
                        unit=unit,
                        price=cost_price
                    ).order_by('id')

                    if existing_items.exists():
                        matched_item = existing_items.first()
                        matched_item.quantity += quantity
                        matched_item.save()
                    else:
                        matched_item = Item.objects.create(
                            company=invoice.company,
                            item_name=item_name,
                            item_code=f"{item_name[:3].upper()}_{Item.objects.count() + 1}",
                            quantity=quantity,
                            unit=unit,
                            description="Auto-created from Purchase Invoice",
                            tax_type=None,
                            tax=None,
                            price=cost_price,
                            selling_price=round(cost_price * 1.1, 2)
                        )

                    # STEP 2.2: Create PurchaseInvoiceItem
                    PurchaseInvoiceItem.objects.create(
                        invoice=invoice,
                        item_name=item_name,
                        unit=unit,
                        quantity=quantity,
                        cost_price=cost_price,
                        line_total=line_total
                    )

                # STEP 3: Update invoice main fields
                invoice.supplier_name = data.get('supplier_name', invoice.supplier_name)
                invoice.invoice_number = data.get('invoice_number', invoice.invoice_number)
                invoice.total_price = round(total_price, 2)
                invoice.save(update_fields=["supplier_name", "invoice_number", "total_price"])

                return Response(PurchaseInvoiceSerializer(invoice).data, status=200)

        except ValidationError as ve:
            return Response({"error": str(ve)}, status=400)
        except Exception as e:
            return Response({"error": f"Unexpected error: {str(e)}"}, status=500)

class ListPurchaseInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="purchase-invoice"

    def get(self, request,company_id):
        invoices = PurchaseInvoice.objects.filter(company_id=company_id,deleted=False)
        serializer = PurchaseInvoiceSerializer(invoices, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

class DeletePurchaseInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="purchase-invoice"

    def delete(self, request,company_id, pk):
        try:
            invoice = PurchaseInvoice.objects.get(pk=pk,company_id=company_id,deleted=False)
            items = invoice.items.all()

            for pi_item in items:
                from company.models import Item
                matching_items = Item.objects.filter(
                    item_name=pi_item.item_name,
                    company=invoice.company,
                    unit=pi_item.unit
                )
                for item in matching_items:
                    item.quantity -= pi_item.quantity
                    item.quantity = max(0, item.quantity)
                    item.save()

            invoice.deleted = True
            invoice.save(update_fields=['deleted'])

            return Response({"detail": "Purchase invoice soft-deleted"}, status=status.HTTP_200_OK)

        except PurchaseInvoice.DoesNotExist:
            return Response({"error": "Invoice not found or already deleted"}, status=status.HTTP_404_NOT_FOUND)

class RetrievePurchaseInvoiceView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="purchase-invoice"

    def get(self, request, company_id,pk):
        try:
            invoice = PurchaseInvoice.objects.get(pk=pk,company_id=company_id,deleted=False)
            serializer = PurchaseInvoiceSerializer(invoice)
            return Response(serializer.data, status=200)
        except PurchaseInvoice.DoesNotExist:
            return Response({"error": "Purchase invoice not found"}, status=404)


            return Response({"error": str(e)}, status=400)







#+=========================================================================================================================
#============================             PAYMENT_IN_OUT                              ======================================================================
#======================================================================================================== 




class PaymentInView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="payment"
    def post(self, request,company_id):
        data = request.data
        amount = float(data.get('amount'))
        invoice_ids = data.get('invoice_ids', [])
        company_id = data.get('company_id')  # ✅ required now
        bank_id = data.get('bank_id')        # optional if cash
        payment_mode = int(data.get('payment_mode', 1))  # 0=cheque, 1=cash, 2=bank transfer
        
        original_amount = amount
        # Validate required fields
        if not amount or not invoice_ids or not company_id:
            return Response({"error": "amount, invoice_ids, and company_id are required."}, status=400)

        try:
            company = Company.objects.get(id=company_id)
        except Company.DoesNotExist:
            return Response({"error": "Company not found."}, status=404)

        bank = None
        if payment_mode in [0, 2]:
            if not bank_id:
                return Response({"error": "bank_id is required for cheque or bank transfer."}, status=400)
            try:
                bank = Bank.objects.get(id=bank_id, company=company)
            except Bank.DoesNotExist:
                return Response({"error": "Bank does not belong to the provided company."}, status=400)

        with transaction.atomic():
            for invoice_id in invoice_ids:
                try:
                    invoice = SalesInvoice.objects.select_for_update().get(id=invoice_id, company=company)
                except SalesInvoice.DoesNotExist:
                    continue

                remaining = invoice.total_price - invoice.received_amt
                if remaining <= 0:
                    continue

                if amount >= remaining:
                    invoice.received_amt += remaining
                    invoice.payment_status_id = 3  # Fully paid
                    amount -= remaining
                else:
                    invoice.received_amt += amount
                    invoice.payment_status_id = 2  # Partially paid
                    amount = 0

                invoice.save(update_fields=['received_amt', 'payment_status'])

                if amount <= 0:
                    break

            description = f"Payment received for invoices: {invoice_ids}"

            if payment_mode in [0, 2]:  # Cheque or Bank Transfer
                BankTransaction.objects.create(
                    bank=bank,
                    company=company,
                    amount=original_amount,
                    transaction_type='credit',
                    description=description
                )
            elif payment_mode == 1:  # Cash
                CashLedger.objects.create(
                    company=company,
                    amount=original_amount,
                    transaction_type='inflow',
                    description=description
                )

        return Response({"status": 200, "message": "Payment applied successfully."})

class PaymentOutView(APIView):
    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="payment"
    def post(self, request,company_id):
        data = request.data
        amount = float(data.get('amount'))
        invoice_ids = data.get('invoice_ids', [])
        bank_id = data.get('bank_id')  # Required only if payment_mode is cheque/bank transfer
        payment_mode = int(data.get('payment_mode', 1))  # Default to cash
        original_amount = amount
        if not amount or not invoice_ids:
            return Response({"error": "amount and invoice_ids are required."}, status=400)

        if payment_mode in [0, 2] and not bank_id:
            return Response({"error": "bank_id is required for cheque or bank transfer."}, status=400)

        try:
            bank = Bank.objects.get(id=bank_id) if bank_id else None
        except Bank.DoesNotExist:
            return Response({"error": "Bank not found."}, status=404)

        with transaction.atomic():
            for invoice_id in invoice_ids:
                try:
                    invoice = PurchaseInvoice.objects.select_for_update().get(company_id=company_id,id=invoice_id)
                except PurchaseInvoice.DoesNotExist:
                    continue

                remaining = invoice.total_price - invoice.paid_amt
                if remaining <= 0:
                    continue

                if amount >= remaining:
                    invoice.paid_amt += remaining
                    invoice.payment_status_id = 3  # Fully paid
                    amount -= remaining
                else:
                    invoice.paid_amt += amount
                    invoice.payment_status_id = 2  # Partially paid
                    amount = 0

                invoice.save(update_fields=['paid_amt', 'payment_status'])

                if amount <= 0:
                    break

            description = f"Payment made for purchase invoices: {invoice_ids}"
            company = invoice.company  # assuming all invoices are from same company

            if payment_mode in [0, 2]:  # Cheque or Bank Transfer
                BankTransaction.objects.create(
                    bank=bank,
                    company=bank.company,
                    amount=original_amount,
                    transaction_type='debit',
                    description=description
                )
            elif payment_mode == 1:  # Cash
                CashLedger.objects.create(
                    company=company,
                    amount=original_amount,
                    transaction_type='outflow',
                    description=description
                )

        return Response({"status": 200, "message": "Payment applied and balance adjusted successfully."})

    

#+=========================================================================================================================
#============================             CREATE_EXCEL                              ======================================================================
#================================================================================================================= 



class InvoiceReportExportView(APIView):

    permission_classes = [IsAuthenticated,OwnerOrEmployee]
    module_name ="excel"
    is_post_as_get=True

    def post(self, request,company_id):
        company_id = request.data.get('company_id')
        invoice_type = request.data.get('invoice_type')  # "sales" or "purchase"
        payment_status = request.data.get('payment_status')  # Optional

        if not company_id or invoice_type not in ["sales", "purchase"]:
            return JsonResponse({"error": "company_id and valid invoice_type are required"}, status=400)

        try:
            company = Company.objects.get(id=company_id)
        except Company.DoesNotExist:
            return JsonResponse({"error": "Invalid company_id"}, status=404)

        if invoice_type == "sales":
            invoices = SalesInvoice.objects.filter(company=company, deleted=False)
        else:
            invoices = PurchaseInvoice.objects.filter(company=company, deleted=False)

        if payment_status is not None:
            invoices = invoices.filter(payment_status_id=payment_status)

        if not invoices.exists():
            return JsonResponse({"error": "No invoices found."}, status=404)

        # Prepare Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = f"{invoice_type.capitalize()} Invoices"

        # Define headers
        headers = [
            "S.No", "Invoice Number", "Date", "Client/Supplier Name",
            "Total Amount", "Paid/Received", "Pending Amount", "Status"
        ]
        ws.append(headers)

        # Styling headers
        header_font = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="4F81BD")
        alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = fill
            cell.alignment = alignment
            cell.border = border
            col_letter = get_column_letter(col_num)
            ws.column_dimensions[col_letter].width = 20

        # Populate rows
        for idx, invoice in enumerate(invoices, start=1):
            if invoice_type == "sales":
                name = invoice.client.client_name
                paid = invoice.received_amt
            else:
                name = invoice.supplier_name
                paid = invoice.paid_amt

            row = [
                idx,
                invoice.invoice_number,
                invoice.invoice_date.strftime("%Y-%m-%d"),
                name,
                invoice.total_price,
                paid,
                round(invoice.total_price - paid, 2),
                invoice.payment_status.label
            ]
            ws.append(row)
            
            last_row = ws.max_row

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=last_row, column=col_num)
                cell.alignment = alignment
                cell.border = border


        # File path
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        filename = f"{invoice_type}_invoice_report_{timestamp}.xlsx"
        file_path = os.path.join(settings.MEDIA_ROOT, filename)
        os.makedirs(settings.MEDIA_ROOT, exist_ok=True)
        wb.save(file_path)

        file_url = f"{request.build_absolute_uri(settings.MEDIA_URL)}{filename}"
        return JsonResponse({"download_url": file_url}, status=200)
    


#+=========================================================================================================================
#============================             Job Role                               ======================================================================
#================================================================================================================= 

class CreateJobRoleWithPermissionsView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "jobrole"

    def post(self, request,company_id):
        data = request.data.copy()
        permissions = data.pop('permissions', [])

        # ✅ Deduplicate module names
        seen_modules = set()
        deduped_permissions = []
        for perm in permissions:
            module = perm['module_name'].strip().lower()
            if module not in seen_modules:
                seen_modules.add(module)
                deduped_permissions.append(perm)

        # ✅ Validate ownership of company
        owner = getattr(request.user, 'owner_profile', None)
        if not owner:
            return Response({"error": "Owner not found."}, status=status.HTTP_400_BAD_REQUEST)

        companies = owner.companies.all()
        company_id = data.get('company')
        if not company_id:
            return Response({"error": "Company ID is required."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            company = companies.get(id=company_id)
        except Company.DoesNotExist:
            return Response({"error": "Invalid company for this user."}, status=status.HTTP_403_FORBIDDEN)

        # ✅ Inject verified company back into data
        data['company'] = company.id

        # ✅ Save JobRole
        serializer = JobRoleCreateSerializer(data=data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        job_role = serializer.save()

        duplicate_modules = []
        for perm in deduped_permissions:
            module_name = perm['module_name'].strip().lower()
            _, created = ModulePermission.objects.get_or_create(
                job_role=job_role,
                company=company,
                module_name=module_name,
                defaults={
                    "can_view": perm.get('can_view', False),
                    "can_create": perm.get('can_create', False),
                    "can_edit": perm.get('can_edit', False),
                    "can_delete": perm.get('can_delete', False),
                    "can_view_specific": perm.get('can_view_specific', False),
                    "can_get_using_post": perm.get('can_get_using_post', False),
                }
            )
            if not created:
                duplicate_modules.append(module_name)

        return Response({
            "message": "Job role created successfully.",
            "job_role_id": job_role.id,
            "name": job_role.name,
            "skipped_modules": duplicate_modules
        }, status=status.HTTP_201_CREATED)
class JobRoleListView(APIView):
    permission_classes=[IsAuthenticated,OwnerOrEmployee]
    module_name ="jobrole"
    def get(self,request,company_id):
        job_roles = JobRole.objects.filter(company_id=company_id,is_deleted=False)
        serializer = JobRoleDetailSerializer(job_roles, many=True)
        return Response(serializer.data)

class JobRoleDetailView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "jobrole"

    def get(self, request,company_id, job_role_id):
        try:
            job_role = JobRole.objects.get(id=job_role_id,company_id=company_id, is_deleted=False)
        except JobRole.DoesNotExist:
            return Response({"error": "Job role not found"}, status=404)

        if hasattr(request.user, 'owner_profile'):
            company_ids = request.user.owner_profile.companies.values_list("id", flat=True)
            if job_role.company.id not in company_ids:
                return Response({"error": "Unauthorized access to job role"}, status=403)
        else:
            employee = Employee.objects.get(user=request.user, is_deleted=False)
            if job_role.company.id != employee.company.id:
                return Response({"error": "Unauthorized access to job role"}, status=403)

        serializer = JobRoleDetailSerializer(job_role)
        return Response(serializer.data)


class JobRoleUpdateView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "jobrole"

    def put(self, request,company_id, job_role_id):
        job_role = get_object_or_404(JobRole, id=job_role_id, is_deleted=False)

        # Company access check for owner or employee
        if hasattr(request.user, 'owner_profile'):
            owner_company_ids = request.user.owner_profile.companies.values_list('id', flat=True)
            if job_role.company.id not in owner_company_ids:
                return Response({"error": "Access denied: Not your company."}, status=status.HTTP_403_FORBIDDEN)
        else:
            if job_role.company.id != request.user.employee_master.company.id:
                return Response({"error": "Access denied: Not your company."}, status=status.HTTP_403_FORBIDDEN)

        # Proceed with update
        data = request.data

        new_name = data.get("name")
        if not new_name:
            return Response({"error": "Name is required."}, status=status.HTTP_400_BAD_REQUEST)

        if JobRole.objects.filter(name=new_name, company=job_role.company).exclude(id=job_role_id).exists():
            return Response({"error": "Job role name must be unique."}, status=status.HTTP_400_BAD_REQUEST)

        job_role.name = new_name
        job_role.save()

        permissions_data = data.get("permissions")
        if permissions_data is None:
            return Response({"error": "Permissions are required."}, status=status.HTTP_400_BAD_REQUEST)

        # Clear existing permissions
        ModulePermission.objects.filter(job_role=job_role).delete()

        # Create new permissions
        for perm in permissions_data:
            module_name = perm.get("module_name")
            if not module_name:
                continue

            internal_name = ALIAS_TO_MODULE_MAP.get(module_name.strip().lower())
            if not internal_name:
                return Response({"error": f"Invalid module name alias: {module_name}"}, status=status.HTTP_400_BAD_REQUEST)

            ModulePermission.objects.create(
                job_role=job_role,
                company=job_role.company,  # ✅ Required to avoid IntegrityError
                module_name=internal_name,
                can_view=perm.get("can_view", False),
                can_create=perm.get("can_create", False),
                can_edit=perm.get("can_edit", False),
                can_delete=perm.get("can_delete", False),
                can_view_specific=perm.get("can_view_specific", False),
                can_get_using_post=perm.get("can_get_using_post", False)
            )

        return Response({"message": "Job role updated successfully."}, status=status.HTTP_200_OK)

class JobRoleDeleteView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "jobrole"

    def delete(self, request,company_id, job_role_id):
        job_role = get_object_or_404(JobRole, id=job_role_id, is_deleted=False)

        # Ensure same company
        if job_role.company.id != request.user.employee_master.company.id:
            return Response({"error": "Access denied: Not your company."}, status=status.HTTP_403_FORBIDDEN)

        job_role.is_deleted = True
        job_role.save()
        return Response({"message": "Job role deleted (soft) successfully."}, status=status.HTTP_200_OK)
 


#+=========================================================================================================================
#============================             EEMPLOYEE                              ======================================================================
#================================================================================================================= 



class CreateEmployeeView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "employees"

    def post(self, request):
        data = request.data
        username = data.get('username')
        email = data.get('email')
        password = data.get('password')
        phone_number = data.get('phone_number')
        company_mappings = data.get('company_mappings', [])  # List of {company_id, job_role_id}

        # Validate required fields
        if not all([username, email, password, phone_number]) or not company_mappings:
            return Response({"error": "All fields and at least one company mapping required"}, status=400)

        # Validate password
        try:
            validate_password(password)
        except ValidationError as e:
            return Response({"error": e.messages}, status=400)

        # Check existing user
        if User.objects.filter(username=username).exists():
            return Response({"error": "Username exists"}, status=400)
        if User.objects.filter(email=email).exists():
            return Response({"error": "Email exists"}, status=400)

        # Create user and employee
        user = User.objects.create_user(username=username, email=email, password=password)
        employee = Employee.objects.create(user=user, phone_number=phone_number)

        # Create company mappings
        created_mappings = []
        for mapping in company_mappings:
            company_id = mapping.get('company_id')
            job_role_id = mapping.get('job_role_id')
            
            if not company_id or not job_role_id:
                continue
                
            try:
                company = Company.objects.get(id=company_id)
                job_role = JobRole.objects.get(id=job_role_id, company=company)
                
                # Create mapping
                EmployeeCompanyMap.objects.create(
                    employee=employee,
                    company=company,
                    job_role=job_role,
                    is_active=True
                )
                created_mappings.append({
                    "company": company.company_name,
                    "job_role": job_role.name
                })
                
            except (Company.DoesNotExist, JobRole.DoesNotExist):
                # Skip invalid mappings
                continue

        return Response({
            "message": "Employee created with company mappings",
            "employee_id": employee.id,
            "username": user.username,
            "company_mappings": created_mappings
        }, status=201)
class AddEmployeeToCompanyView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "employees"

    def post(self, request, employee_id):
        data = request.data
        company_mappings = data.get('company_mappings', [])  # List of {company_id, job_role_id}

        if not company_mappings:
            return Response({"error": "At least one company mapping required"}, status=400)

        try:
            employee = Employee.objects.get(id=employee_id)
        except Employee.DoesNotExist:
            return Response({"error": "Employee not found"}, status=404)

        # Create new mappings
        created_mappings = []
        for mapping in company_mappings:
            company_id = mapping.get('company_id')
            job_role_id = mapping.get('job_role_id')
            
            if not company_id or not job_role_id:
                continue
                
            try:
                company = Company.objects.get(id=company_id)
                job_role = JobRole.objects.get(id=job_role_id, company=company)
                
                # Create new mapping
                EmployeeCompanyMap.objects.create(
                    employee=employee,
                    company=company,
                    job_role=job_role,
                    is_active=True
                )
                created_mappings.append({
                    "company": company.company_name,
                    "job_role": job_role.name
                })
                
            except (Company.DoesNotExist, JobRole.DoesNotExist):
                # Skip invalid mappings
                continue

        return Response({
            "message": "Employee added to new companies",
            "employee_id": employee.id,
            "added_mappings": created_mappings
        }, status=201)
    

class EmployeeDetailView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]

    def get(self, request, employee_id):
        try:
            employee = Employee.objects.get(id=employee_id, deleted=False)
        except Employee.DoesNotExist:
            return Response({"error": "Employee not found."}, status=404)

        # Get all active mappings for this employee
        mappings = EmployeeCompanyMap.objects.filter(
            employee=employee,
            is_active=True
        ).select_related('company', 'job_role')

        # Serialize employee details
        data = {
            "id": employee.id,
            "username": employee.user.username,
            "email": employee.user.email,
            "phone_number": employee.phone_number,
            "mappings": [
                {
                    #"mapping_id": mapping.id,
                    "company_id": mapping.company.id,
                    "company_name": mapping.company.company_name,
                    "job_role_id": mapping.job_role.id,
                    "job_role_name": mapping.job_role.name
                }
                for mapping in mappings
            ]
        }

        return Response(data, status=200)
class EmployeeListView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "employees"

    def get(self, request):
        # Owners see all employees in their companies
        if hasattr(request.user, 'owner_profile'):
            companies = request.user.owner_profile.companies.all()
            mappings = EmployeeCompanyMap.objects.filter(
                company__in=companies,
                is_active=True,
                employee__deleted=False
            ).select_related('employee', 'company', 'job_role', 'employee__user')
        
        # Superusers see all employees
        elif request.user.is_superuser:
            mappings = EmployeeCompanyMap.objects.filter(
                is_active=True,
                employee__deleted=False
            ).select_related('employee', 'company', 'job_role', 'employee__user')
        
        # Employees see only themselves
        elif hasattr(request.user, 'employee'):
            mappings = EmployeeCompanyMap.objects.filter(
                employee=request.user.employee,
                is_active=True
            ).select_related('employee', 'company', 'job_role', 'employee__user')
        
        else:
            return Response({"error": "Permission denied."}, status=403)

        # Group by employee
        employees = {}
        for mapping in mappings:
            emp = mapping.employee
            if emp.id not in employees:
                employees[emp.id] = {
                    "id": emp.id,
                    "username": emp.user.username,
                    "email": emp.user.email,
                    "phone_number": emp.phone_number,
                    "mappings": []
                }
            
            employees[emp.id]["mappings"].append({
                #"mapping_id": mapping.id,
                "company_id": mapping.company.id,
                "company_name": mapping.company.company_name,
                "job_role_id": mapping.job_role.id,
                "job_role_name": mapping.job_role.name
            })

        return Response(list(employees.values()), status=200)
class EmployeeUpdateView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "employees"

    def put(self, request, employee_id):
        try:
            employee = Employee.objects.get(id=employee_id, deleted=False)
        except Employee.DoesNotExist:
            return Response({"error": "Employee not found."}, status=404)
        
        data = request.data
        
        # Update basic info
        if 'phone_number' in data:
            employee.phone_number = data['phone_number']
            employee.save()
        
        # Update specific mapping (e.g., change job role in a company)
        if 'mapping_id' in data and 'job_role_id' in data:
            try:
                mapping = EmployeeCompanyMap.objects.get(
                    id=data['mapping_id'],
                    employee=employee
                )
                mapping.job_role = JobRole.objects.get(id=data['job_role_id'])
                mapping.save()
            except (EmployeeCompanyMap.DoesNotExist, JobRole.DoesNotExist):
                return Response({"error": "Invalid mapping or job role"}, status=400)
        
        return Response({"message": "Employee updated successfully"})
class EmployeeDeleteView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "employees"

    def delete(self, request, employee_id):
        try:
            employee = Employee.objects.get(id=employee_id, deleted=False)
        except Employee.DoesNotExist:
            return Response({"error": "Employee not found."}, status=404)
        
        # Soft delete employee and deactivate all mappings
        employee.deleted = True
        employee.save()
        
        EmployeeCompanyMap.objects.filter(employee=employee).update(is_active=False)
        
        return Response({"message": "Employee deleted successfully"})

class RemoveEmployeeRoleView(APIView):
    permission_classes = [IsAuthenticated, OwnerOrEmployee]
    module_name = "employees"

    def delete(self, request, mapping_id):
        try:
            mapping = EmployeeCompanyMap.objects.get(id=mapping_id)
            mapping.is_active = False
            mapping.save()
            return Response({"message": "Employee removed from company/role"})
            
        except EmployeeCompanyMap.DoesNotExist:
            return Response({"error": "Mapping not found"}, status=404)